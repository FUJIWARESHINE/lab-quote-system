# -*- coding: utf-8 -*-
"""Excel 生成：客户报价单、项目库模板与导入导出。"""
from __future__ import annotations

import math
import os
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import db

# ------------------------------------------------------------------ 样式
FONT = "微软雅黑"
BLUE = "1F4E79"
BLUE_LIGHT = "D9E2F3"
GREY = "F2F2F2"

THIN = Side(style="thin", color="9DA9BF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_TEXT = {
    "draft": "草稿",
    "sent": "已发送",
    "confirmed": "已确认",
    "rejected": "已拒绝",
}


def _int_to_cn(n: int) -> str:
    """整数转中文大写（支持到万亿）。"""
    if n == 0:
        return "零"
    digits = "零壹贰叁肆伍陆柒捌玖"
    units = ["", "拾", "佰", "仟"]
    sections = ["", "万", "亿", "万亿"]
    groups, m = [], n
    while m > 0:
        groups.append(m % 10000)
        m //= 10000
    result, prev_len = "", 4
    for i, g in enumerate(groups):  # 由低位到高位，每 4 位一段
        if g == 0:
            if result and not result.startswith("零"):
                result = "零" + result
            prev_len = 0
            continue
        part, zero = "", False
        for j in range(4):
            d = (g // (10 ** j)) % 10
            if d == 0:
                zero = True
            else:
                if zero and part:
                    part = "零" + part
                part = digits[d] + units[j] + part
                zero = False
        if result and prev_len < 4 and not result.startswith("零"):
            result = "零" + result
        result = part + sections[i] + result
        prev_len = len(str(g))
    return result


def _upper_rmb(amount) -> str:
    """人民币金额转中文大写。"""
    try:
        num = round(float(amount) + 0.0, 2)
    except Exception:
        return ""
    neg = num < 0
    num = abs(num)
    if num >= 1e16:
        return "金额过大"
    integer = int(num)
    cents = int(round((num - integer) * 100))
    if integer == 0 and cents == 0:
        return "零元整"
    text = _int_to_cn(integer) + "元"
    if cents == 0:
        text += "整"
    else:
        jiao, fen = cents // 10, cents % 10
        text += ("零" if jiao == 0 and fen else "")
        text += (("零壹贰叁肆伍陆柒捌玖"[jiao] + "角") if jiao else "")
        text += (("零壹贰叁肆伍陆柒捌玖"[fen] + "分") if fen else "")
    return ("负" + text) if neg else text


def _wrap_lines(text: str, units_per_line: float) -> int:
    """估算文本在给定列宽（字符数）下换行后的实际行数（中文按约 2 个字符宽计）。"""
    total = 0
    for paragraph in str(text or "").split("\n"):
        width = sum(1.95 if ord(ch) > 0x2E80 else 1.05 for ch in paragraph)
        total += max(1, math.ceil(width / max(10.0, units_per_line)))
    return total


def _setup_page(ws, columns: int) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9  # A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "第 &P 页 / 共 &N 页"
    ws.oddFooter.center.size = 9
    ws.oddFooter.center.font = "宋体"


def _fill(ws, cell_range, value=None, font=None, fill=None, align=None, border=None):
    for row in ws[cell_range]:
        for cell in row:
            if value is not None:
                cell.value = value
            if font:
                cell.font = font
            if fill:
                cell.fill = fill
            if align:
                cell.alignment = align
            if border:
                cell.border = border


# ------------------------------------------------------------------ 报价单
def _build_quote_workbook(quote: dict, settings: dict) -> Workbook:
    """按报价单内容构建 Excel 工作簿（不落盘）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "报价单"
    cols = {"A": 6, "B": 14, "C": 23, "D": 19, "E": 7, "F": 10, "G": 9,
            "H": 12, "I": 9, "J": 9, "K": 11, "L": 13, "M": 20}
    for letter, width in cols.items():
        ws.column_dimensions[letter].width = width
    total_cols = 13
    last = get_column_letter(total_cols)

    _setup_page(ws, total_cols)

    company = settings.get("company_name") or ""
    row = 1
    ws.merge_cells(f"A{row}:{last}{row}")
    c = ws[f"A{row}"]
    c.value = company
    c.font = Font(name=FONT, size=17, bold=True, color=BLUE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 32

    row += 1
    ws.merge_cells(f"A{row}:{last}{row}")
    c = ws[f"A{row}"]
    c.value = "检 测 报 价 单"
    c.font = Font(name=FONT, size=15, bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 26

    row += 1
    ws.merge_cells(f"A{row}:{last}{row}")
    c = ws[f"A{row}"]
    sub = settings.get("company_addr") or ""
    if settings.get("company_phone"):
        sub = f"{sub}    电话：{settings['company_phone']}".strip("    ")
    c.value = sub
    c.font = Font(name=FONT, size=9, color="666666")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 16

    row += 1
    info = [
        ("客户单位：", quote.get("customer_name") or "", "报价单号：", quote.get("quote_no") or ""),
        ("联 系 人：", quote.get("contact") or "", "报价日期：", quote.get("quote_date") or ""),
        ("联系电话：", quote.get("phone") or "", "有效期至：", quote.get("expire_date") or ""),
        ("项目地址：", quote.get("address") or "", "报价状态：", STATUS_TEXT.get(quote.get("status", ""), "")),
    ]
    label_font = Font(name=FONT, size=10, bold=True)
    value_font = Font(name=FONT, size=10)
    # 标签跨列合并，避免「客户单位：」等文字被窄列裁掉（列宽保持不变）
    left_units = sum(cols[get_column_letter(i)] for i in (3, 4, 5)) - 2
    right_units = sum(cols[get_column_letter(i)] for i in range(8, total_cols + 1)) - 2
    for left_label, left_value, right_label, right_value in info:
        ws.merge_cells(f"A{row}:B{row}")
        ws[f"A{row}"] = left_label
        ws[f"A{row}"].font = label_font
        ws[f"A{row}"].alignment = Alignment(horizontal="right", vertical="center")
        ws.merge_cells(f"C{row}:E{row}")
        ws[f"C{row}"] = left_value
        ws[f"C{row}"].font = value_font
        ws[f"C{row}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.merge_cells(f"F{row}:G{row}")
        ws[f"F{row}"] = right_label
        ws[f"F{row}"].font = label_font
        ws[f"F{row}"].alignment = Alignment(horizontal="right", vertical="center")
        ws.merge_cells(f"H{row}:{last}{row}")
        ws[f"H{row}"] = right_value
        ws[f"H{row}"].font = value_font
        ws[f"H{row}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        # 客户名称 / 地址较长时自动加高，保证内容完整显示
        lines = max(_wrap_lines(left_value, left_units), _wrap_lines(right_value, right_units), 1)
        ws.row_dimensions[row].height = max(19, 15 * lines + 4)
        row += 1

    ws.merge_cells(f"A{row}:{last}{row}")
    ws[f"A{row}"] = quote.get("title") or "检测服务报价"
    ws[f"A{row}"].font = Font(name=FONT, size=11, bold=True, color=BLUE)
    ws[f"A{row}"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 20
    row += 1

    headers = ["序号", "项目编号", "检测项目", "规格 / 检测依据", "单位", "分组数",
               "每组\n样本数", "检测指标数\n/时间点", "重复\n次数", "合计\n数量",
               "单价(元)", "金额(元)", "备注"]

    def write_head(r: int) -> None:
        for idx, name in enumerate(headers, start=1):
            cell = ws.cell(row=r, column=idx, value=name)
            cell.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER
        ws.row_dimensions[r].height = 30

    items = quote.get("items") or []
    if not items:
        write_head(row)
        row += 1
        ws.merge_cells(f"A{row}:{last}{row}")
        ws[f"A{row}"] = "暂无报价明细"
        ws[f"A{row}"].alignment = Alignment(horizontal="center")
        ws[f"A{row}"].font = Font(name=FONT, size=10, color="999999")
        ws[f"A{row}"].border = BORDER
        row += 1
    else:
        # 按“分组/部分”分段输出，每段含标题、表头、明细与小计
        sections, grouped = [], {}
        for it in items:
            name = (it.get("section") or "").strip() or "报价明细"
            if name not in grouped:
                grouped[name] = []
                sections.append(name)
            grouped[name].append(it)

        seq = 0
        for sec_index, sec_name in enumerate(sections, start=1):
            ws.merge_cells(f"A{row}:{last}{row}")
            ws[f"A{row}"] = sec_name
            ws[f"A{row}"].font = Font(name=FONT, size=11, bold=True, color=BLUE)
            ws[f"A{row}"].alignment = Alignment(horizontal="left", vertical="center")
            for col in range(1, total_cols + 1):
                ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=BLUE_LIGHT)
                ws.cell(row=row, column=col).border = BORDER
            ws.row_dimensions[row].height = 22
            row += 1
            write_head(row)
            row += 1

            section_total = 0.0
            for local_idx, it in enumerate(grouped[sec_name], start=1):
                seq += 1
                qty = float(it.get("quantity") or 0)
                price = float(it.get("price") or 0)
                dr = float(it.get("discount_rate") or 0)
                amount = round(price * qty * (1 - dr / 100), 2)
                section_total += amount
                note = it.get("note") or ""
                values = [
                    seq, it.get("code") or "", it.get("name") or "", it.get("spec") or "",
                    it.get("unit") or "项", float(it.get("group_count") or 1),
                    float(it.get("samples") or 0) or "",
                    float(it.get("indicators") or 0) or "", float(it.get("repeats") or 0) or "",
                    qty, price, amount, note,
                ]
                zebra = local_idx % 2 == 0
                for col, value in enumerate(values, start=1):
                    cell = ws.cell(row=row, column=col, value=value)
                    cell.font = Font(name=FONT, size=10)
                    cell.border = BORDER
                    if zebra:
                        cell.fill = PatternFill("solid", fgColor=GREY)
                    if col in (1, 2, 3, 4, 5, 13):
                        # 序号 / 项目编号 / 检测项目 / 规格 / 单位 / 备注 居中
                        cell.alignment = Alignment(horizontal="center", vertical="center",
                                                   wrap_text=True)
                    elif col in (6, 7, 8, 9, 10):
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.number_format = "0"  # 数量统计，不显示小数点
                    else:
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = "#,##0.00"
                if dr:
                    ws.cell(row=row, column=12).font = Font(name=FONT, size=10, color="C00000")
                if note:
                    # 备注用红色，字号与「检测项目」等正文一致，并按内容自动加高
                    ws.cell(row=row, column=13).font = Font(name=FONT, size=10, color="C00000")
                    note_lines = _wrap_lines(note, cols["M"] - 2)
                    ws.row_dimensions[row].height = max(20, 15 * note_lines + 5)
                else:
                    ws.row_dimensions[row].height = 20
                row += 1

            ws.merge_cells(f"A{row}:K{row}")
            ws[f"A{row}"] = f"{sec_name}  小计"
            ws[f"A{row}"].font = Font(name=FONT, size=10, bold=True)
            ws[f"A{row}"].alignment = Alignment(horizontal="right", vertical="center")
            st = ws.cell(row=row, column=12, value=round(section_total, 2))
            st.font = Font(name=FONT, size=10, bold=True, color=BLUE)
            st.number_format = "#,##0.00"
            st.alignment = Alignment(horizontal="right", vertical="center")
            for col in range(1, total_cols + 1):
                ws.cell(row=row, column=col).border = BORDER
            ws.row_dimensions[row].height = 20
            row += 1

    subtotal = float(quote.get("subtotal") or 0)
    discount_amount = float(quote.get("discount_amount") or 0)
    tax_amount = float(quote.get("tax_amount") or 0)
    total = float(quote.get("total") or 0)

    ws.merge_cells(f"A{row}:K{row}")
    ws[f"A{row}"] = f"金额合计（大写）：{_upper_rmb(total)}"
    ws[f"A{row}"].font = Font(name=FONT, size=11, bold=True)
    ws[f"A{row}"].alignment = Alignment(horizontal="right", vertical="center")
    tc = ws.cell(row=row, column=12, value=total)
    tc.font = Font(name=FONT, size=11, bold=True, color="C00000")
    tc.number_format = "#,##0.00"
    tc.alignment = Alignment(horizontal="right", vertical="center")
    for col in range(1, total_cols + 1):
        ws.cell(row=row, column=col).border = BORDER
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=BLUE_LIGHT)
    ws.row_dimensions[row].height = 24
    row += 1

    note = (
        f"小计：¥ {subtotal:,.2f}    折扣（{float(quote.get('discount_rate') or 0):g}%）："
        f"-¥ {discount_amount:,.2f}    税额（{float(quote.get('tax_rate') or 0):g}%）：¥ {tax_amount:,.2f}"
    )
    ws.merge_cells(f"A{row}:{last}{row}")
    ws[f"A{row}"] = note
    ws[f"A{row}"].font = Font(name=FONT, size=9, color="555555")
    ws[f"A{row}"].alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[row].height = 17
    row += 1

    terms = (quote.get("terms") or settings.get("terms") or "").replace("{days}", settings.get("valid_days") or "30")
    if terms:
        ws.merge_cells(f"A{row}:B{row}")
        ws[f"A{row}"] = "报价说明："
        ws[f"A{row}"].font = Font(name=FONT, size=10, bold=True)
        ws[f"A{row}"].alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(f"C{row}:{last}{row}")
        ws[f"C{row}"] = terms
        ws[f"C{row}"].font = Font(name=FONT, size=9)
        ws[f"C{row}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        terms_units = sum(cols[get_column_letter(i)] for i in range(3, total_cols + 1)) - 4
        text_lines = max(4, _wrap_lines(terms, terms_units))
        ws.row_dimensions[row].height = 15 * text_lines + 6
        row += 1

    if quote.get("remark"):
        ws.merge_cells(f"A{row}:B{row}")
        ws[f"A{row}"] = "备注："
        ws[f"A{row}"].font = Font(name=FONT, size=10, bold=True)
        ws[f"A{row}"].alignment = Alignment(horizontal="center", vertical="top")
        ws.merge_cells(f"C{row}:{last}{row}")
        ws[f"C{row}"] = quote.get("remark")
        ws[f"C{row}"].font = Font(name=FONT, size=9)
        ws[f"C{row}"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = 26
        row += 1

    row += 1
    ws.merge_cells(f"A{row}:D{row}")
    ws[f"A{row}"] = "报价方（盖章）：" + (company or "")
    ws[f"A{row}"].font = Font(name=FONT, size=10, bold=True)
    ws[f"A{row}"].alignment = Alignment(horizontal="left", vertical="bottom")
    ws.merge_cells(f"E{row}:{last}{row}")
    ws[f"E{row}"] = "客户确认（签字/盖章）："
    ws[f"E{row}"].font = Font(name=FONT, size=10, bold=True)
    ws[f"E{row}"].alignment = Alignment(horizontal="left", vertical="bottom")
    ws.row_dimensions[row].height = 22
    row += 1
    ws.merge_cells(f"A{row}:D{row}")
    ws[f"A{row}"] = "日期：          年        月        日"
    ws[f"A{row}"].font = Font(name=FONT, size=10)
    ws[f"A{row}"].alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(f"E{row}:{last}{row}")
    ws[f"E{row}"] = "日期：          年        月        日"
    ws[f"E{row}"].font = Font(name=FONT, size=10)
    ws[f"E{row}"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 22

    return wb


def quote_filename(quote: dict) -> str:
    """报价单 Excel 的标准文件名。"""
    safe_customer = re.sub(r'[\\/:*?"<>|]', "_", (quote.get("customer_name") or "客户"))[:20]
    no = quote.get("quote_no") or datetime.now().strftime("%Y%m%d")
    return f"报价单_{safe_customer}_{no}.xlsx"


def quote_workbook(quote: dict, settings: dict) -> Workbook:
    """按报价单内容构建 Excel 工作簿对象（内存中，不落盘）。"""
    return _build_quote_workbook(quote, settings)


def quote_bytes(quote: dict, settings: dict) -> tuple:
    """在内存中生成报价单 Excel，返回 (bytes, filename)。"""
    import io

    wb = _build_quote_workbook(quote, settings)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), quote_filename(quote)


def export_quote(quote: dict, settings: dict) -> str:
    """生成报价单 Excel 并保存到服务器导出目录，返回文件路径。"""
    wb = _build_quote_workbook(quote, settings)
    save_dir = db.resolve_export_dir(settings.get("export_dir"))
    os.makedirs(save_dir, exist_ok=True)
    path = os.path.join(save_dir, quote_filename(quote))
    wb.save(path)
    return path


# ------------------------------------------------------------------ 项目库
ITEM_HEADERS = [
    "项目编号", "项目名称", "分类", "规格/检测依据", "单位",
    "单价", "趋势价", "趋势价说明", "成本价", "备注", "启用", "收藏",
]
ITEM_WIDTHS = [14, 30, 14, 22, 8, 12, 12, 22, 12, 20, 8, 8]


def parse_price_text(text) -> tuple:
    """从 '60元/对'、'100-200元/样' 等文本中解析。

    返回 (数值或 None, 单位, 原文)：无法判定为单一价格时数值为 None。
    """
    raw = "" if text is None else str(text).strip()
    if not raw:
        return None, "", ""
    if isinstance(text, (int, float)):
        return float(text), "", ""
    multi = re.search(r"[-~～—]|；|;|或|起|以上|以内|<|>", raw)
    m = re.search(r"\d+(?:\.\d+)?", raw)
    value = float(m.group()) if m else None
    unit = ""
    mu = re.search(r"元[/\s]*([\u4e00-\u9fa5A-Za-z/·]+)$", raw)
    if mu:
        unit = mu.group(1).strip()
    if multi:
        return None, unit, raw
    return value, unit, raw


def export_items(rows: list) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "项目库"
    for i, w in enumerate(ITEM_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col, name in enumerate(ITEM_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.freeze_panes = "A2"
    for r, item in enumerate(rows, start=2):
        values = [
            item.get("code") or "", item.get("name") or "", item.get("category_name") or "",
            item.get("spec") or "", item.get("unit") or "项", float(item.get("price") or 0),
            float(item.get("price2") or 0) or "", item.get("price2_text") or "",
            float(item.get("cost") or 0), item.get("note") or "", "是" if item.get("enabled", 1) else "否",
            "是" if item.get("starred") else "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            if col in (6, 7, 9):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[r].height = 19
    save_dir = db.get_data_dir()
    path = os.path.join(save_dir, f"项目库导出_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    wb.save(path)
    return path


def export_items_template() -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "项目库模板"
    for i, w in enumerate(ITEM_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col, name in enumerate(ITEM_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    samples = [
        ["LHBZ-003", "蛋白质测定", "理化检测", "GB 5009.5", "项", 180, 200, "", 60, "示例数据可删除", "是", "是"],
        ["WSW-003", "霉菌酵母菌", "微生物检测", "GB 4789.15", "项", 160, "", "100-200元/样（区间，需手填）", 50, "", "是", ""],
    ]
    for r, row_values in enumerate(samples, start=2):
        for col, value in enumerate(row_values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            if col in (6, 7, 9):
                cell.number_format = "#,##0.00"
    ws["M1"] = "填写说明：单价＝常规价；趋势价填数字，区间或条件价请留空并填到「趋势价说明」；收藏列填「是」即置顶"
    ws["M1"].font = Font(name=FONT, size=9, color="888888")
    dv = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("K2:K1000")  # 启用
    dv.add("L2:L1000")  # 收藏
    path = os.path.join(db.get_data_dir(), "项目库导入模板.xlsx")
    wb.save(path)
    return path


def _norm_header(text) -> str:
    return re.sub(r"[\s（）()【】/、\-_]*", "", str(text or ""))


def import_items(path: str) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"created": 0, "updated": 0}
    header_map = {}
    for idx, name in enumerate(rows[0]):
        key = _norm_header(name)
        for std in ITEM_HEADERS:
            if key and key in _norm_header(std):
                header_map[std] = idx
                break
    if "项目名称" not in header_map:
        return {"created": 0, "updated": 0, "error": "未识别到「项目名称」列，请使用标准模板"}

    categories = {c["name"]: c["id"] for c in db.list_categories()}
    cat_name_by_id = {v: k for k, v in categories.items()}
    existing = {}
    with db.tx() as conn:
        for r in conn.execute("SELECT id, code, name, category_id FROM items"):
            cn = cat_name_by_id.get(r["category_id"] or 0, "")
            if r["code"]:
                existing[r["code"]] = r["id"]
            existing.setdefault(r["name"], r["id"])
            existing[f"{cn}|{r['name']}"] = r["id"]

    created = updated = 0
    for raw in rows[1:]:
        def cell_of(std):
            idx = header_map.get(std)
            if idx is None or idx >= len(raw):
                return ""
            return raw[idx]

        name = str(cell_of("项目名称") or "").strip()
        if not name:
            continue
        code = str(cell_of("项目编号") or "").strip()
        cat_name = str(cell_of("分类") or "").strip()
        if cat_name and cat_name not in categories:
            categories[cat_name] = db.create_category({"name": cat_name})
        p_value, p_unit, _ = parse_price_text(cell_of("单价"))
        t_raw = cell_of("趋势价")
        t_value, _, t_orig = parse_price_text(t_raw)
        t_text = str(cell_of("趋势价说明") or "").strip()
        if t_value is None and t_raw not in (None, ""):
            t_text = t_text or t_orig
        payload = {
            "code": code,
            "name": name,
            "category_id": categories.get(cat_name, 0),
            "spec": str(cell_of("规格/检测依据") or ""),
            "unit": str(cell_of("单位") or "项").strip() or (p_unit or "项"),
            "price": p_value if p_value is not None else 0.0,
            "price2": t_value if t_value is not None else 0.0,
            "price2_text": t_text,
            "cost": float(cell_of("成本价") or 0),
            "note": str(cell_of("备注") or ""),
            "enabled": 0 if str(cell_of("启用") or "是").strip() in ("否", "0", "false", "False", "禁用") else 1,
            "starred": 1 if str(cell_of("收藏") or "").strip() in ("是", "1", "true", "True", "★", "收藏") else 0,
        }
        key = f"{cat_name}|{name}" if cat_name else name
        iid = existing.get(key) or (existing.get(code) if code else None) or existing.get(name)
        if iid:
            db.update_item(iid, payload)
            updated += 1
        else:
            iid = db.create_item(payload)
            existing[key] = iid
            existing.setdefault(name, iid)
            created += 1
    return {"created": created, "updated": updated}
