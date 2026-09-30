# -*- coding: utf-8 -*-
"""生成内置初始数据：为整理好的项目加上正式编号，写入 seed.db 并输出带编号的 Excel。"""
import os
import shutil
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(ROOT, "app"))
import db  # noqa: E402
from export_excel import ITEM_HEADERS, ITEM_WIDTHS, parse_price_text  # noqa: E402

SRC = "F:/Download/报价单（新）.xlsx"
SEED = os.path.join(ROOT, "seed.db")
OUT_XLSX = os.path.join(ROOT, "导入项目_科信研测报价单.xlsx")
PREFIX = "KX"

CAT_CODE = {
    "RT-qPCR": "RTQ", "WB": "WB", "ELISA": "ELISA", "培养": "PY", "转染": "ZZ",
    "流式": "LS", "功能": "GN", "取材": "QC", "切片制备": "QP", "染色": "RS",
    "拍照": "PZ", "分析": "FX", "动物代购": "DW", "饲养": "SY", "操作": "CZ", "模型": "MX",
}

THIN = Side(style="thin", color="9DA9BF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BLUE = "1F4E79"


def cell_text(ws, r, c):
    v = ws.cell(row=r, column=c).value
    return "" if v is None else str(v).replace("\n", " ").strip()


def parse_src():
    wb = load_workbook(SRC, data_only=True)
    ws = wb.worksheets[0]
    rows, platform, experiment = [], "", ""
    for r in range(4, ws.max_row + 1):
        b, c = cell_text(ws, r, 2), cell_text(ws, r, 3)
        name, price_raw, trend_raw = cell_text(ws, r, 4), cell_text(ws, r, 5), cell_text(ws, r, 6)
        if b:
            platform = b
        if c:
            experiment = c
        if not name or not price_raw:
            continue
        p_value, p_unit, _ = parse_price_text(price_raw)
        t_value, _, t_orig = parse_price_text(trend_raw)
        t_text = t_orig if (t_value is None and trend_raw) else ""
        rows.append({
            "name": name,
            "category": experiment or platform or "未分类",
            "spec": platform,
            "unit": p_unit or "项",
            "price": p_value if p_value is not None else 0,
            "price2": t_value if t_value is not None else 0,
            "price2_text": t_text,
            "note": "" if p_value is not None else f"常规价待确认：{price_raw}",
        })
    return rows


def main():
    rows = parse_src()
    counter = {}
    for item in rows:
        code = CAT_CODE.get(item["category"], "QT")
        counter[code] = counter.get(code, 0) + 1
        item["code"] = f"{PREFIX}-{code}{counter[code]:03d}"

    # 1) 重建数据库（去掉示例数据，仅保留正式项目）
    path = db.get_db_path()
    for suffix in ("", "-wal", "-shm"):
        if os.path.exists(path + suffix):
            os.remove(path + suffix)
    db.init_db()
    with db.tx() as conn:
        for table in ("quote_items", "quotes", "items", "customers", "categories"):
            conn.execute(f"DELETE FROM {table}")

    cats = {}
    for item in rows:
        name = item["category"]
        if name not in cats:
            cats[name] = db.create_category({"name": name})
        payload = dict(item)
        payload["category_id"] = cats[name]
        payload["enabled"] = 1
        payload["starred"] = 0
        db.create_item(payload)

    # 2) 导出带编号的 Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "项目库导入"
    for i, w in enumerate(ITEM_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col, name in enumerate(ITEM_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    for r, item in enumerate(rows, start=2):
        values = [
            item["code"], item["name"], item["category"], item["spec"], item["unit"],
            item["price"], item["price2"] or "", item["price2_text"], "", item["note"], "是", "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = Font(name="微软雅黑", size=10)
            cell.border = BORDER
            if col in (6, 7, 9):
                cell.number_format = "#,##0.00"
            cell.alignment = Alignment(vertical="center", wrap_text=(col in (2, 4, 8, 10)))
    ws.freeze_panes = "A2"
    wb.save(OUT_XLSX)

    # 3) 生成内置数据模块 app/seed_data.py（编译进 exe，最可靠）
    cat_order = []
    for item in rows:
        if item["category"] not in cat_order:
            cat_order.append(item["category"])
    lines = [
        "# -*- coding: utf-8 -*-",
        '"""内置初始项目数据（由 make_seed.py 自动生成，请勿手工修改）。"""',
        "",
        "SEED_CATEGORIES = [",
    ]
    lines += [f"    {name!r}," for name in cat_order]
    lines.append("]")
    lines.append("")
    lines.append("# 项目编号, 名称, 分类, 规格/平台, 单位, 常规价, 趋势价, 趋势价说明, 备注")
    lines.append("SEED_ITEMS = [")
    for item in rows:
        lines.append(
            f"    [{item['code']!r}, {item['name']!r}, {item['category']!r}, {item['spec']!r}, "
            f"{item['unit']!r}, {item['price']}, {item['price2']}, "
            f"{item['price2_text']!r}, {item['note']!r}],"
        )
    lines.append("]")
    lines.append("")
    with open(os.path.join(ROOT, "app", "seed_data.py"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    total = len(rows)
    trend = sum(1 for x in rows if x["price2"])
    print("项目数:", total, "| 有趋势价:", trend, "| 分类:", len(cats))
    print("编号示例:", [x["code"] for x in rows[:3]], "...", rows[-1]["code"])
    print("Excel:", OUT_XLSX, os.path.getsize(OUT_XLSX))
    mod = os.path.join(ROOT, "app", "seed_data.py")
    print("seed_data.py:", mod, os.path.getsize(mod) // 1024, "KB")


if __name__ == "__main__":
    main()
